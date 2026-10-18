"""Build a vertical (9:16) Douyin promo deck for dsh-term-glossary.

Every slide carries the spoken script in its notes, so it can be recorded or
edited straight from PowerPoint / Keynote / 剪映.

usage: python build-pptx.py   → promo/douyin.pptx
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

BG = RGBColor(0x0D, 0x0F, 0x14)
PANEL = RGBColor(0x16, 0x1A, 0x22)
PANEL_LINE = RGBColor(0x2B, 0x33, 0x42)
BUBBLE = RGBColor(0x17, 0x1C, 0x26)
FG = RGBColor(0xF3, 0xF5, 0xF9)
DIM = RGBColor(0x9A, 0xA4, 0xB5)
BRAND = RGBColor(0x7D, 0x92, 0xFF)
BRAND_DEEP = RGBColor(0x4D, 0x6B, 0xFE)
CHIP = RGBColor(0x12, 0x17, 0x22)
CHIP_HOT = RGBColor(0x16, 0x20, 0x3A)

FONT = "Microsoft YaHei"

prs = Presentation()
prs.slide_width = Inches(11.25)
prs.slide_height = Inches(20)
BLANK = prs.slide_layouts[6]

PAD = Inches(0.9)
W = prs.slide_width
INNER = W - PAD * 2


def new_slide():
    slide = prs.slides.add_slide(BLANK)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG
    return slide


def text(slide, x, y, w, h, lines, spacing=1.25):
    """lines: list of (text, size_pt, bold, color) — one paragraph each."""
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for index, (value, size, bold, color) in enumerate(lines):
        para = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        para.alignment = PP_ALIGN.LEFT
        para.line_spacing = spacing
        run = para.add_run()
        run.text = value
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = FONT
    return box


def box(slide, x, y, w, h, fill, line=PANEL_LINE, radius=0.08):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1.2)
    shape.adjustments[0] = radius
    shape.shadow.inherit = False
    return shape


def kicker(slide, value):
    text(slide, PAD, Inches(1.55), INNER, Inches(0.5), [(value, 22, True, BRAND)])


def title(slide, value, size=60, top=2.05):
    text(slide, PAD, Inches(top), INNER, Inches(4.2), [(value, size, True, FG)], spacing=1.16)


def note(slide, value):
    slide.notes_slide.notes_text_frame.text = value


# ── 0 · 封面 ────────────────────────────────────────────────────────────
slide = new_slide()
box(slide, PAD, Inches(6.4), Inches(2.6), Inches(0.14), BRAND_DEEP, line=BRAND_DEEP)
text(slide, PAD, Inches(6.9), INNER, Inches(0.6), [("DeepSeek Harness 插件", 26, True, BRAND)])
text(slide, PAD, Inches(7.7), INNER, Inches(3.4), [("术语速查", 132, True, FG)], spacing=1.0)
text(slide, PAD, Inches(11.0), INNER, Inches(1.6), [("选中术语，就地解释", 48, True, FG)], spacing=1.2)
text(slide, PAD, Inches(12.6), INNER, Inches(1.2), [("基本含义 · 用途 · 用法，顺手看 token 成本", 30, False, DIM)])
note(
    slide,
    "封面（约 1.5 秒）。画面只留标题，制造悬念。\n"
    "口播：'给你看个我刚做的小插件——读文档时遇到不认识的词，选中它就能问。'",
)

# ── 1 · 痛点 ────────────────────────────────────────────────────────────
slide = new_slide()
kicker(slide, "读文档 / 看教程时")
title(slide, "又遇到\n看不懂的术语？")
text(slide, PAD, Inches(6.15), INNER, Inches(1.6), [("每次都要新开一个对话去问，阅读的节奏就断了。", 32, False, DIM)])
box(slide, PAD, Inches(8.3), INNER, Inches(3.2), BUBBLE)
text(slide, PAD + Inches(0.45), Inches(8.7), INNER - Inches(0.9), Inches(2.4),
     [("会话需要保证「幂等性」，重复提交不会产生副作用……", 32, False, RGBColor(0xE6, 0xEA, 0xF3))])
box(slide, PAD, Inches(11.9), INNER, Inches(1.7), RGBColor(0x1A, 0x24, 0x40), RGBColor(0x2B, 0x39, 0x60))
text(slide, PAD + Inches(0.45), Inches(12.25), INNER - Inches(0.9), Inches(1.0),
     [("（幂等性是啥？再开个对话问一下…）", 30, False, DIM)])
note(
    slide,
    "钩子（约 3 秒）。手指或光标划过『幂等性』这个高亮词。\n"
    "口播：'读文档遇到不认识的术语，你是不是又要新开一个对话去问？一打断，思路就没了。'",
)

# ── 2 · 一步 ────────────────────────────────────────────────────────────
slide = new_slide()
kicker(slide, "现在只要一步")
title(slide, "选中它，\n点一下「搜索」")
text(slide, PAD, Inches(6.15), INNER, Inches(1.4), [("解释就出现在原地——不用切窗口，也不用新开对话。", 32, False, DIM)])
box(slide, PAD, Inches(8.0), INNER, Inches(3.4), BUBBLE)
text(slide, PAD + Inches(0.45), Inches(8.5), INNER - Inches(0.9), Inches(2.5),
     [("会话需要保证 ", 32, False, RGBColor(0xE6, 0xEA, 0xF3)),
      ("幂等性", 32, True, FG),
      ("，重复提交不会产生副作用……", 32, False, RGBColor(0xE6, 0xEA, 0xF3))])
box(slide, Inches(6.35), Inches(11.9), Inches(2.95), Inches(0.95), BUBBLE, RGBColor(0x3A, 0x42, 0x56), radius=0.5)
text(slide, Inches(6.6), Inches(12.12), Inches(2.6), Inches(0.6), [("搜索", 28, True, FG)])
note(
    slide,
    "演示核心（约 4 秒）。录屏时鼠标从『幂等性』上松手 → 按钮浮出 → 点击。\n"
    "口播：'现在不用了。选中它，旁边会冒出一个搜索按钮，点一下。'",
)

# ── 3 · 就地解释 ────────────────────────────────────────────────────────
slide = new_slide()
kicker(slide, "基本含义 · 用途 · 用法")
title(slide, "一次说清", size=54)
box(slide, PAD, Inches(4.6), INNER, Inches(6.6), PANEL)
text(slide, PAD + Inches(0.5), Inches(5.0), INNER - Inches(1.0), Inches(0.7), [("幂等性", 40, True, FG)])
text(slide, PAD + Inches(0.5), Inches(5.85), INNER - Inches(1.0), Inches(4.8), [
    ("同一个操作执行一次和多次，对系统状态的影响相同。", 30, True, FG),
    ("", 12, False, DIM),
    ("用途：网络超时重试、用户连点提交，都不会造成重复扣款。", 28, False, RGBColor(0xE7, 0xEB, 0xF3)),
    ("用法：接口带唯一 requestId，服务端记住已处理的请求。", 28, False, RGBColor(0xE7, 0xEB, 0xF3)),
], spacing=1.45)
box(slide, PAD, Inches(11.55), INNER, Inches(0.8), RGBColor(0x12, 0x17, 0x22), PANEL_LINE, radius=0.2)
text(slide, PAD + Inches(0.4), Inches(11.74), INNER - Inches(0.8), Inches(0.5),
     [("token：输入 204 · 输出 330 · 合计 534 · deepseek-flash", 24, False, DIM)])
note(
    slide,
    "效果（约 5 秒）。先露『正在思考…』再出答案，最后一行 token 用高亮。\n"
    "口播：'它会就地给你基本含义、用途、用法，还会告诉你这次解释花掉多少 token。'",
)

# ── 4 · 不打断 ─────────────────────────────────────────────────────────
slide = new_slide()
kicker(slide, "关键是")
title(slide, "它不打断\n你手上的活")
for index, line in enumerate([
    "不写进会话记录，不占上下文轮次",
    "正在跑的任务、正在写的代码，都不受影响",
    "顺手告诉你：这次解释花了多少 token",
]):
    top = Inches(7.0 + index * 1.75)
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, PAD, top + Inches(0.18), Inches(0.18), Inches(0.18))
    dot.fill.solid()
    dot.fill.fore_color.rgb = BRAND_DEEP
    dot.line.fill.background()
    dot.shadow.inherit = False
    text(slide, PAD + Inches(0.5), top, INNER - Inches(0.5), Inches(1.2), [(line, 34, False, FG)])
note(
    slide,
    "价值点（约 3.5 秒）。三条依次出现。\n"
    "口播：'重点是它不占你对话的上下文，也不打断正在跑的任务；顺便还能看到成本。'",
)

# ── 5 · 懂领域 ─────────────────────────────────────────────────────────
slide = new_slide()
kicker(slide, "而且它")
title(slide, "按你正在聊的领域\n来回答", size=54)
text(slide, PAD, Inches(6.0), INNER, Inches(1.6),
     [("它读的是你这段对话本身，同一个词在医学、法律里不会被硬套成技术术语。", 30, False, DIM)])
chips = ["计算机", "医学", "法律", "金融", "化学", "物理"]
x, y = PAD, Inches(8.2)
for index, value in enumerate(chips):
    width = Inches(2.05 if len(value) <= 2 else 2.6)
    if x + width > W - PAD:
        x, y = PAD, y + Inches(1.25)
    box(slide, x, y, width, Inches(0.9), CHIP_HOT if index == 0 else CHIP, BRAND_DEEP if index == 0 else PANEL_LINE, radius=0.5)
    text(slide, x + Inches(0.28), y + Inches(0.22), width - Inches(0.4), Inches(0.5), [(value, 26, False, FG)])
    x += width + Inches(0.25)
note(
    slide,
    "差异化（约 4 秒）。词可以逐个弹出。\n"
    "口播：'它还会根据你正在聊的领域来回答，学医的、学法律的，不会被解释成程序员那套。'",
)

# ── 6 · 安装 ───────────────────────────────────────────────────────────
slide = new_slide()
kicker(slide, "30 秒装好")
title(slide, "三步，不用命令行", size=54)
for index, line in enumerate([
    "① 设置 → 插件 → 市场",
    "② 搜「术语速查」或 dsh-term-glossary",
    "③ 点安装",
]):
    top = Inches(5.6 + index * 1.9)
    text(slide, PAD, top, INNER, Inches(1.2), [(line, 34, False, FG)])
text(slide, PAD, Inches(11.3), INNER, Inches(1.6),
     [("也可以在对话里直接说：", 28, False, DIM),
      ("「帮我安装插件 dsh-term-glossary」", 28, True, BRAND)])
note(
    slide,
    "安装（约 4 秒）。可以插一段真实的插件市场录屏。\n"
    "口播：'装起来也就三步：设置里进插件市场，搜术语速查，点安装。'",
)

# ── 7 · CTA ────────────────────────────────────────────────────────────
slide = new_slide()
kicker(slide, "DeepSeek Harness 插件")
text(slide, PAD, Inches(3.4), INNER, Inches(2.2), [("术语速查", 96, True, FG)], spacing=1.0)
text(slide, PAD, Inches(5.9), INNER, Inches(1.4),
     [("选中术语就地解释 · 不打断对话 · 顺手看成本", 28, False, DIM)])
box(slide, PAD, Inches(7.8), INNER, Inches(3.2), BUBBLE)
text(slide, PAD + Inches(0.5), Inches(8.2), INNER - Inches(1.0), Inches(2.6), [
    ("GitHub：github.com/ashz555/dsh-term-glossary", 26, False, RGBColor(0xE7, 0xEB, 0xF3)),
    ("市场：设置 → 插件 → 市场，搜「术语速查」", 26, False, RGBColor(0xE7, 0xEB, 0xF3)),
    ("命令：dsh plugin add dsh-term-glossary", 26, True, FG),
], spacing=1.5)
note(
    slide,
    "收尾（约 3 秒）。停住 2 秒方便截图/打字幕。\n"
    "口播：'插件叫术语速查，商店里搜名字就能装，开源免费的，链接放评论区。'",
)

out = Path(__file__).with_name("douyin.pptx")
prs.save(out)
print(f"saved {out} ({len(prs.slides.__iter__.__self__._sldIdLst)} slides, {prs.slide_width}x{prs.slide_height} EMU)")
